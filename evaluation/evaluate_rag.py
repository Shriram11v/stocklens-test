"""
Runs a gold Q&A set (from an .xlsx built by build_gold_qa.py) against
BM25 / Dense / Hybrid retrievers from src/, computes retrieval metrics
(Hit@k, MRR, nDCG@k), and writes per-question results to JSON for later
manual generation scoring (build_scoring_template.py).

SETUP
    python evaluate_rag.py Gold_QA_Set_CBA_DEV.xlsx --k 5
    python evaluate_rag.py Gold_QA_Set_CBA_DEV.xlsx --k 5 --retrievers bm25 dense
    python evaluate_rag.py Gold_QA_Set_CBA_DEV.xlsx --k 5 --generate   # also runs generator.py (needs Ollama)
    python evaluate_rag.py Gold_QA_Set_CBA_TEST.xlsx --k 5 --generate  # only once, at the end
    python evaluate_rag.py Gold_QA_Set_CBA_DEV.xlsx --project-root /path/to/repo

Each retriever is built ONCE (BM25 index / dense FAISS index + embedding
model load) and then reused for every question — rebuilding per-question
would reload the embedding model and re-encode all ~1,000 chunks every
single time, which is both slow and pointless since the corpus doesn't
change between questions.

WHAT COUNTS AS A "HIT"
A retrieved chunk counts as correct if its "page" field (the pipeline's own
per-chunk page number — see the gold set's Corpus Notes tab) is IN the gold
"Page(s)" list for that question. Multi-page gold answers (e.g. "16, 17")
count a hit if the retrieved chunk matches ANY of those pages.

COMPANY FILTER
All CBA gold questions are evaluated with company="CBA" so NAB chunks in
the shared index can't accidentally count as noise or false hits.
"""

import argparse
import json
import math
import sys
from pathlib import Path

import openpyxl


# ------------------------------------------------------------------
# 0. Locate the repo root (the folder containing src/ and processed_data/)
#    and make src/ importable, matching how hybrid_retriever.py imports
#    bm25_retriever / dense_retriever directly (not as a package).
#
#    This file doesn't need to sit at repo root — it walks upward from
#    its own location (e.g. an evaluation/ subfolder) looking for a
#    directory that contains both src/ and processed_data/.
# ------------------------------------------------------------------

def find_project_root(start, max_up=5):
    current = Path(start).resolve()
    for _ in range(max_up + 1):
        if (current / "src").is_dir() and (current / "processed_data").is_dir():
            return current
        if current.parent == current:  # reached filesystem root
            break
        current = current.parent
    return None


def setup_src_path(explicit_root=None):
    if explicit_root:
        root = Path(explicit_root).resolve()
        if not (root / "src").is_dir():
            print(f"! --project-root {root} has no src/ folder — check the path.", file=sys.stderr)
    else:
        root = find_project_root(Path(__file__).resolve().parent)
        if root is None:
            print(
                "! Couldn't find a folder containing both src/ and processed_data/ by walking up "
                f"from {Path(__file__).resolve().parent}. Pass --project-root explicitly.",
                file=sys.stderr,
            )
            return None
    src_dir = root / "src"
    sys.path.insert(0, str(src_dir))
    bm25_file = src_dir / "bm25_retriever.py"
    if not bm25_file.exists():
        print(
            f"! Detected project root as {root}, and added {src_dir} to the import path, "
            f"but {bm25_file} doesn't exist there. Contents of {src_dir}:",
            file=sys.stderr,
        )
        if src_dir.is_dir():
            for p in sorted(src_dir.iterdir()):
                print(f"    {p.name}", file=sys.stderr)
        else:
            print(f"    (this folder doesn't exist)", file=sys.stderr)
    return root


# ------------------------------------------------------------------
# 1. LOAD THE GOLD SET
# ------------------------------------------------------------------

def load_gold_set(xlsx_path):
    """Reads the Dev or Test sheet from a gold-set workbook into a list of dicts."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]  # first sheet is the Q&A sheet; others are notes/glossary

    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if row and row[0] == "ID":
            header_row_idx = i
            break
    if header_row_idx is None:
        raise ValueError("Couldn't find header row (looking for 'ID' in column A)")

    headers = [c.value for c in ws[header_row_idx]]
    records = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row[0] is None:
            continue
        rec = dict(zip(headers, row))
        pages_raw = str(rec.get("Page(s)", ""))
        rec["gold_pages"] = [int(p.strip()) for p in pages_raw.split(",") if p.strip().isdigit()]
        records.append(rec)
    return records


# ------------------------------------------------------------------
# 2. BUILD RETRIEVERS (once each, reused across all questions)
# ------------------------------------------------------------------

def build_retrievers(which):
    """
    Returns {name: callable(question, k) -> list of result dicts} for the
    requested retriever names, built once each against the real chunk data.
    """
    from bm25_retriever import BM25Retriever, load_chunks as load_chunks_bm25, CHUNKS_FILE

    print(f"Loading chunks from {CHUNKS_FILE} ...")
    chunks = load_chunks_bm25(CHUNKS_FILE)
    print(f"Loaded {len(chunks)} usable chunks.\n")

    retrievers = {}

    if "bm25" in which:
        print("Building BM25 retriever...")
        bm25 = BM25Retriever(chunks)
        retrievers["bm25"] = lambda q, k: bm25.search(query=q, top_k=k, company="CBA")

    if "dense" in which:
        print("Building Dense retriever (loads embedding model + FAISS index)...")
        from dense_retriever import DenseRetriever
        dense = DenseRetriever(chunks)
        retrievers["dense"] = lambda q, k: dense.search(query=q, top_k=k, company="CBA")

    if "hybrid" in which:
        print("Building Hybrid retriever (this also builds its own internal BM25 + Dense)...")
        from hybrid_retriever import HybridRetriever
        hybrid = HybridRetriever(chunks)
        retrievers["hybrid"] = lambda q, k: hybrid.search(query=q, top_k=k, company="CBA")

    print()
    return retrievers


def load_generator():
    """
    Lazily imports generate_answer from generator.py. Returns None (with a
    warning) if Ollama isn't installed/running — generation is optional;
    retrieval evaluation still works without it.
    """
    try:
        from generator import generate_answer
        return generate_answer
    except Exception as e:
        print(f"! Could not load generator.py ({e}). Skipping generation.", file=sys.stderr)
        return None


# ------------------------------------------------------------------
# 3. RETRIEVAL METRICS — Hit@k, MRR, nDCG@k (mirrors the Walert reproduction)
# ------------------------------------------------------------------

def rank_of_first_hit(retrieved_pages, gold_pages):
    for i, page in enumerate(retrieved_pages, start=1):
        if page in gold_pages:
            return i
    return None


def hit_at_k(retrieved_pages, gold_pages, k):
    return int(any(p in gold_pages for p in retrieved_pages[:k]))


def reciprocal_rank(retrieved_pages, gold_pages):
    rank = rank_of_first_hit(retrieved_pages, gold_pages)
    return 1.0 / rank if rank else 0.0


def ndcg_at_k(retrieved_pages, gold_pages, k):
    """Binary relevance nDCG@k: relevant chunk = 1, else 0."""
    dcg = 0.0
    for i, page in enumerate(retrieved_pages[:k], start=1):
        rel = 1 if page in gold_pages else 0
        dcg += rel / math.log2(i + 1)
    ideal_hits = min(k, len(gold_pages)) if gold_pages else 0
    idcg = sum(1 / math.log2(i + 1) for i in range(1, ideal_hits + 1))
    return dcg / idcg if idcg > 0 else 0.0


# ------------------------------------------------------------------
# 4. RUN EVALUATION
# ------------------------------------------------------------------

def evaluate(gold_records, retriever_name, retriever_fn, k=5, generate_fn=None, gen_top_k=3):
    rows = []
    for rec in gold_records:
        question = rec["Question"]
        gold_pages = rec["gold_pages"]
        try:
            results = retriever_fn(question, k)
        except Exception as e:
            print(f"  ! retrieval failed for ID {rec['ID']}: {e}", file=sys.stderr)
            results = []

        retrieved_pages = [r["page"] for r in results]
        retrieved_texts = [r.get("text", "") for r in results]

        rank = rank_of_first_hit(retrieved_pages, gold_pages)
        row = {
            "id": rec["ID"],
            "question": question,
            "gold_pages": gold_pages,
            "retrieved_pages": retrieved_pages,
            "retrieved_texts": retrieved_texts,
            "hit_at_k": hit_at_k(retrieved_pages, gold_pages, k),
            "rr": reciprocal_rank(retrieved_pages, gold_pages),
            "ndcg_at_k": ndcg_at_k(retrieved_pages, gold_pages, k),
            "first_hit_rank": rank,
        }

        if generate_fn is not None:
            # generator.py is built around top-3 evidence (its own DEFAULT_TOP_K),
            # so feed it the same slice it would get in normal use, not all k.
            gen_results = results[:gen_top_k]
            try:
                row["generated_answer"] = generate_fn(query=question, results=gen_results)
            except Exception as e:
                print(f"  ! generation failed for ID {rec['ID']}: {e}", file=sys.stderr)
                row["generated_answer"] = None

        rows.append(row)

    n = len(rows)
    summary = {
        "retriever": retriever_name,
        "n_questions": n,
        f"Hit@{k}": sum(r["hit_at_k"] for r in rows) / n if n else 0,
        "MRR": sum(r["rr"] for r in rows) / n if n else 0,
        f"nDCG@{k}": sum(r["ndcg_at_k"] for r in rows) / n if n else 0,
    }
    return summary, rows


def main():
    parser = argparse.ArgumentParser(description="Evaluate FinTrace retrievers against a gold Q&A set.")
    parser.add_argument("gold_xlsx", help="Path to a Gold_QA_Set_*.xlsx file (Dev or Test)")
    parser.add_argument("--k", type=int, default=5, help="Cutoff for Hit@k / nDCG@k (default 5)")
    parser.add_argument("--retrievers", nargs="+", default=["bm25", "dense", "hybrid"],
                         choices=["bm25", "dense", "hybrid"],
                         help="Which retrievers to evaluate (default: all three)")
    parser.add_argument("--out", default="eval_results.json", help="Where to write per-question results")
    parser.add_argument("--generate", action="store_true",
                         help="Also run generator.py on retrieved evidence and save the generated answer "
                              "(requires Ollama installed and running locally with the model pulled)")
    parser.add_argument("--gen-retriever", default="hybrid", choices=["bm25", "dense", "hybrid"],
                         help="Which retriever's results to feed the generator (default: hybrid, "
                              "matching generator.py's own default setup)")
    parser.add_argument("--gen-top-k", type=int, default=3,
                         help="How many top results to pass to the generator (default 3, matching "
                              "generator.py's DEFAULT_TOP_K)")
    parser.add_argument("--project-root", default=None,
                         help="Path to the repo root (folder containing src/ and processed_data/). "
                              "Only needed if auto-detection fails for your layout.")
    args = parser.parse_args()

    root = setup_src_path(args.project_root)
    if root is None:
        sys.exit(1)
    print(f"Using project root: {root}\n")

    gold_records = load_gold_set(args.gold_xlsx)
    print(f"Loaded {len(gold_records)} gold questions from {args.gold_xlsx}\n")

    retrievers = build_retrievers(args.retrievers)

    generate_fn = load_generator() if args.generate else None
    if args.generate and generate_fn is not None and args.gen_retriever not in retrievers:
        print(f"! --gen-retriever '{args.gen_retriever}' wasn't built (not in --retrievers); "
              f"generation will be skipped.", file=sys.stderr)
        generate_fn = None

    all_results = {}
    summaries = []
    for name, fn in retrievers.items():
        print(f"Evaluating retriever: {name} ...")
        use_generate = generate_fn if (args.generate and name == args.gen_retriever) else None
        summary, rows = evaluate(gold_records, name, fn, k=args.k,
                                  generate_fn=use_generate, gen_top_k=args.gen_top_k)
        summaries.append(summary)
        all_results[name] = rows

    print("\n=== Summary ===")
    header = f"{'Retriever':<10} {'Hit@'+str(args.k):<10} {'MRR':<10} {'nDCG@'+str(args.k):<10}"
    print(header)
    print("-" * len(header))
    for s in summaries:
        print(f"{s['retriever']:<10} {s[f'Hit@{args.k}']:<10.3f} {s['MRR']:<10.3f} {s[f'nDCG@{args.k}']:<10.3f}")

    with open(args.out, "w") as f:
        json.dump({"summaries": summaries, "per_question": all_results}, f, indent=2)
    print(f"\nFull per-question results written to {args.out}")


if __name__ == "__main__":
    main()
