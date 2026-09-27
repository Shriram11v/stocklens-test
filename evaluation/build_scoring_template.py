"""
Builds a manual scoring template for generation quality: context relevance,
answer faithfulness, answer relevance — each on the same 1-5 scale used in
the RAG evaluation paper the team is following (Iaroshev et al., 2024).

Run: python build_scoring_template.py Gold_QA_Set_CBA_DEV.xlsx eval_results.json output.xlsx
"""
import sys
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def main():
    if len(sys.argv) < 2:
        print("Usage: python build_scoring_template.py <gold.xlsx> [eval_results.json] [output.xlsx]")
        sys.exit(1)
    gold_path = sys.argv[1]
    results_path = sys.argv[2] if len(sys.argv) > 2 else None
    out_path = sys.argv[3] if len(sys.argv) > 3 else "generation_scoring_template.xlsx"

    wb_gold = openpyxl.load_workbook(gold_path, data_only=True)
    ws_gold = wb_gold.worksheets[0]
    header_row_idx = next(i for i, row in enumerate(ws_gold.iter_rows(values_only=True), start=1) if row and row[0] == "ID")
    headers = [c.value for c in ws_gold[header_row_idx]]
    gold_records = []
    for row in ws_gold.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row[0] is None:
            continue
        gold_records.append(dict(zip(headers, row)))

    results_by_id = {}
    if results_path:
        with open(results_path) as f:
            data = json.load(f)
        preferred = "hybrid" if "hybrid" in data.get("per_question", {}) else next(iter(data.get("per_question", {})), None)
        if preferred:
            for row in data["per_question"][preferred]:
                results_by_id[row["id"]] = row

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Generation Scoring"

    FONT = "Arial"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(["Generation Quality Scoring — 1 to 5 scale (see Rubric tab)"])
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws.append([])

    cols = ["ID", "Question", "Gold Answer", "Retrieved Page(s) (actual)", "Retrieved Text (top result)",
            "Generated Answer (paste here)", "Context Relevance (1-5)", "Answer Faithfulness (1-5)",
            "Answer Relevance (1-5)", "Notes"]
    widths = [6, 35, 35, 16, 40, 40, 14, 14, 14, 30]
    header_row = 3
    for i, (name, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=header_row, column=i, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(i)].width = w

    r = header_row + 1
    for rec in gold_records:
        rid = rec["ID"]
        result_row = results_by_id.get(rid, {})
        retrieved_pages = ", ".join(str(p) for p in result_row.get("retrieved_pages", [])) or ""
        retrieved_texts = result_row.get("retrieved_texts", [])
        top_text = (retrieved_texts[0][:300] + "...") if retrieved_texts and len(retrieved_texts[0]) > 300 else (retrieved_texts[0] if retrieved_texts else "")
        generated = result_row.get("generated_answer") or ""
        values = [rid, rec["Question"], rec["Gold Answer"], retrieved_pages, top_text, generated, "", "", "", ""]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = wrap
            cell.border = border
        ws.row_dimensions[r].height = 60
        r += 1

    ws.freeze_panes = f"A{header_row+1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{r-1}"

    # Rubric 
    rubric = wb.create_sheet("Rubric")
    rubric.column_dimensions["A"].width = 12
    rubric.column_dimensions["B"].width = 110
    rubric.append(["Score", "Definition (applies to all three metrics)"])
    for c in (1, 2):
        cell = rubric.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        cell.fill = header_fill
    scale = [
        (1, "The retrieved context is irrelevant to the question, the generated response is inaccurate and inconsistent, and it fails to address the question."),
        (2, "The retrieved context has limited relevance and contains several inaccuracies, the generated response exhibits notable inconsistencies, and it only partially addresses the question."),
        (3, "The retrieved context is somewhat relevant but has some inaccuracies, the generated response is generally consistent with minor errors, and it adequately addresses the question."),
        (4, "The retrieved context is relevant and mostly accurate, the generated response is consistent with minor or no errors, and it effectively addresses the question."),
        (5, "The retrieved context is highly relevant and accurate, the generated response is entirely faithful to the context with no errors, and it thoroughly addresses the question."),
    ]
    for i, (score, definition) in enumerate(scale, start=2):
        rubric.cell(row=i, column=1, value=score).font = Font(name=FONT, bold=True, size=11)
        dcell = rubric.cell(row=i, column=2, value=definition)
        dcell.font = Font(name=FONT, size=11)
        dcell.alignment = wrap
        rubric.row_dimensions[i].height = 45
    rubric.append([])
    rubric.append(["Note", "This rubric is adapted from the RAG evaluation paper the team is following "
                            "(Iaroshev et al., 2024). Score each of the three columns independently — a response "
                            "can have highly relevant context (5) but still be unfaithful to it (e.g. hallucinating "
                            "a figure not actually in the retrieved chunk)."])
    for c in (1, 2):
        rubric.cell(row=rubric.max_row, column=c).font = Font(name=FONT, italic=True, size=10)
        rubric.cell(row=rubric.max_row, column=c).alignment = wrap

    wb.save(out_path)
    print(f"Saved scoring template to {out_path} ({len(gold_records)} questions)")

if __name__ == "__main__":
    main()
